#!/usr/bin/env python3
"""Prefill industrial/dind.xlsx from its local Pages manifest (no network).

python scripts/prefill_industrial_gt.py            # extract + JSON report, no Excel writes
python scripts/prefill_industrial_gt.py --write    # backup, then fill missing GT cells

Rules are based on industrial/html/cadr_examples. These are review candidates,
not verified labels: annotation_status=in_progress, annotator=auto:card-prefill.
No cards found does NOT imply a negative page. Existing values and done/skip
pages are preserved. Repeated runs do not duplicate page/product URL pairs.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
import html
import json
import math
from pathlib import Path
import re
import shutil
from urllib.parse import parse_qsl, unquote, urljoin, urlsplit, urlunsplit

from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ANNOTATOR = "auto:card-prefill"
PRODUCT_FIELDS = (
    "is_product", "name_gt", "price_gt", "product_url_gt", "in_stock_gt",
    "rating_gt", "delivery_time_gt", "region_gt",
)


def clean(value) -> str:
    return re.sub(r"\s+", " ", html.unescape(str(value or ""))).strip()


def blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


QUERY_PARAMETERS = {"q", "text", "query", "search", "searchtext", "keyword"}


def query_from_url(url: str | None) -> str | None:
    """Last nonempty search parameter; preserve encoded literal '+' and '&'."""
    if not url:
        return None
    try:
        pairs = parse_qsl(urlsplit(html.unescape(str(url))).query, keep_blank_values=True)
    except ValueError:
        return None
    for key, value in reversed(pairs):
        if key.casefold() not in QUERY_PARAMETERS:
            continue
        # Some snapshots double-encode UTF-8 percent sequences. Do not call
        # unquote_plus again: a literal %2B in C++ must remain a plus sign.
        for _ in range(2):
            if not re.search(r"%[0-9a-fA-F]{2}", value):
                break
            value = unquote(value)
        value = clean(value)
        if value and "\ufffd" not in value:
            return value
    return None


def recover_query(page: dict, directory: Path) -> dict:
    """Prefer captured page URL; never infer chronology from unrelated links."""
    if not blank(page.get("query")):
        return {"query": page["query"], "query_source": "existing:Pages.query"}
    for key in ("page_url", "source_url"):
        value = query_from_url(page.get(key))
        if value:
            return {"query": value, "query_source": key, "query_url": page[key]}
    path = (directory / page["benchmark_html_file"]).resolve()
    if not path.is_relative_to(directory.resolve()):
        raise ValueError(f"HTML path outside dataset: {path}")
    if not path.is_file():
        return {"query": None, "query_source": "missing_html"}
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for node in soup.select('link[rel="canonical"], meta[property="og:url"]'):
        url = node.get("href") or node.get("content")
        value = query_from_url(url)
        if value:
            return {"query": value, "query_source": "html:page_url", "query_url": url}
    inputs = [clean(n.get("value")) for n in soup.select("input[name][value]")
              if n["name"].casefold() in QUERY_PARAMETERS and clean(n["value"])]
    if inputs and len(set(inputs)) == 1:
        return {"query": inputs[0], "query_source": "html:search_input"}
    base = page.get("page_url") or f"https://{page['source']}/"
    host = (urlsplit(base).hostname or "").removeprefix("www.")
    links = []
    # Only pagination is evidence of the current search. Header promotional
    # links and recommendations can carry completely unrelated ?q= values.
    for node in soup.select('a[rel="next"], a[rel="prev"], [class*="pagination"] a[href], [class*="Pagination"] a[href]'):
        url = urljoin(base, node["href"])
        try:
            if (urlsplit(url).hostname or "").removeprefix("www.") != host:
                continue
        except ValueError:
            continue
        value = query_from_url(url)
        if value:
            links.append((value, url))
    if links and len({v for v, _ in links}) == 1:
        return {"query": links[-1][0], "query_source": "html:pagination_link", "query_url": links[-1][1]}
    # These redirected M.Video snapshots lose ?q= but retain the exact query
    # in their search-page title. Record that this is title evidence, not a URL.
    if page.get("source") == "mvideo.ru" and soup.title:
        match = re.match(r"^(.+?)\s+-\s+купить в интернет магазине", soup.title.get_text(), re.I)
        if match:
            return {"query": clean(match[1]), "query_source": "html:search_title"}
    return {"query": None, "query_source": "ambiguous_links" if links else "not_found",
            "query_candidates": sorted({v for v, _ in links})}


def update_metadata(workbook, directory: Path, *, confirm_empty_pending=False,
                    confirm_unpriced_unavailable=False) -> dict:
    """Update query independently of product extraction, including reviewed rows.

    Confirmation flags encode an explicit human review, never parser guesses.
    A zero price is a user-provided missing-price sentinel and stays untouched.
    """
    query_pages = []
    query_changes = 0
    page_query_column = [c.value for c in workbook["Pages"][1]].index("query") + 1
    for row_num, page in enumerate(records(workbook["Pages"]), 2):
        recovered = recover_query(page, directory)
        query_cell = workbook["Pages"].cell(row_num, page_query_column)
        if not blank(page.get("query")) and query_cell.comment and query_cell.comment.author == ANNOTATOR:
            try:
                previous = json.loads(query_cell.comment.text)
                if previous.get("query") == page["query"]:
                    recovered = previous
            except (ValueError, AttributeError):
                pass
        result = {"html_file": page["benchmark_html_file"], "source": page["source"], **recovered}
        query_pages.append(result)
        if blank(page.get("query")) and recovered["query"]:
            cell = query_cell
            cell.value = recovered["query"]
            cell.data_type = "s"
            cell.comment = Comment(json.dumps(recovered, ensure_ascii=False), ANNOTATOR)
            query_changes += 1
    by_file = {p["html_file"]: p for p in query_pages}
    sheet = workbook["GroundTruth"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    negative_rows = unavailable_rows = 0
    for row_num, row in enumerate(records(sheet), 2):
        recovered = by_file.get(row["html_file"], {})
        if blank(row.get("query")) and recovered.get("query"):
            cell = sheet.cell(row_num, columns["query"], recovered["query"])
            cell.data_type = "s"
            cell.comment = Comment(json.dumps(recovered, ensure_ascii=False), ANNOTATOR)
            query_changes += 1
        note = None
        if (confirm_empty_pending and row.get("annotation_status") == "pending"
                and blank(row.get("is_product"))
                and blank(row.get("name_gt")) and blank(row.get("product_url_gt"))
                and (blank(row.get("price_gt")) or row.get("price_gt") == 0)):
            sheet.cell(row_num, columns["is_product"], 0)
            sheet.cell(row_num, columns["annotation_status"], "done")
            sheet.cell(row_num, columns["annotator"], row.get("annotator") or "user-review")
            note = "USER_REVIEW: page contains no products"
            negative_rows += 1
        elif (confirm_unpriced_unavailable and row.get("is_product") == 1
              and (blank(row.get("price_gt")) or row.get("price_gt") == 0)):
            note = "USER_REVIEW: no price; product no longer sold (0 means missing price)"
            if row.get("in_stock_gt") != 0 or note not in (row.get("notes") or ""):
                sheet.cell(row_num, columns["in_stock_gt"], 0)
                unavailable_rows += 1
        if note and note not in (row.get("notes") or ""):
            sheet.cell(row_num, columns["notes"], "; ".join(filter(None, [row.get("notes"), note])))
    # Update audit query columns even when product extraction changes no rows.
    audit = workbook["GT_Prefill_Audit"] if "GT_Prefill_Audit" in workbook else workbook.create_sheet("GT_Prefill_Audit")
    if audit["A1"].value is None:
        audit.cell(1, 1, "html_file")
    audit_columns = {cell.value: cell.column for cell in audit[1] if cell.value}
    for key in ("query", "query_source"):
        if key not in audit_columns:
            audit_columns[key] = audit.max_column + 1
            audit.cell(1, audit_columns[key], key)
    audit_changed = 0
    # An empty audit contains no stale page list: create just the manifest keys.
    if audit.max_row == 1:
        for p in query_pages:
            audit.append([p["html_file"]])
    for row_num in range(2, audit.max_row + 1):
        recovered = by_file.get(audit.cell(row_num, audit_columns["html_file"]).value, {})
        for key in ("query", "query_source"):
            cell = audit.cell(row_num, audit_columns[key])
            if cell.value != recovered.get(key):
                cell.value = recovered.get(key)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                audit_changed += 1
    audit.auto_filter.ref = audit.dimensions
    return {"query_cells_filled": query_changes, "confirmed_empty_rows": negative_rows,
            "confirmed_unavailable_rows": unavailable_rows, "audit_cells_updated": audit_changed,
            "pages": query_pages}


def number(value) -> float | None:
    """Read one price; don't merge a current and a crossed-out price."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value) if math.isfinite(value) and value >= 0 else None
    text = clean(value)
    match = re.search(r"(?<![\w])-?\d+(?:[ \u00a0\u2009\u202f]\d{3})*(?:\s*[.,]\s*\d{1,2})?", text)
    if not match:
        return None
    price = float(re.sub(r"\s+", "", match.group()).replace(",", "."))
    return price if price >= 0 else None


def canonical_url(value, base: str) -> str | None:
    if blank(value) or str(value).startswith(("#", "javascript:", "mailto:")):
        return None
    parts = urlsplit(urljoin(base, html.unescape(str(value))))
    if parts.scheme not in {"http", "https"} or not parts.netloc:
        return None
    # Same page/product identity as the notebook: ignore query and fragment.
    return urlunsplit((parts.scheme, parts.netloc.lower(), parts.path.rstrip("/"), "", ""))


def first(root: Tag, selectors: str) -> Tag | None:
    # Comma-separated selectors here express priority, not document order.
    for selector in selectors.split(","):
        if selector.strip() == ":self":
            return root
        node = root.select_one(selector.strip())
        if node is not None:
            return node
    return None


def text_of(node: Tag | None) -> str:
    if node is None:
        return ""
    return clean(node.get("content") or node.get_text(" ", strip=True))


@dataclass(frozen=True)
class Rule:
    cards: str
    name: str
    price: str
    link: str = ""
    rating: str = ""
    delivery: str = ""


RULES = {
    "bigam.ru": Rule(".digi-product, .a-product-card", ".digi-product__label, .a-product-card__title", ".digi-product-price-variant, .digi-product__price, .a-price__new, .a-price__current"),
    "chipdip.ru": Rule("tr.with-hover", "a.link", '[id^="price_"]', delivery=".nw"),
    "citilink.ru": Rule('[data-meta-name="SnippetProductVerticalLayout"], [data-meta-name="SnippetProductHorizontalLayout"]', '[data-meta-name="Snippet__title"], a[title][href*="/product/"]', '[class*="MainPriceNumber"]'),
    "cnc.su": Rule(".js-notice-block", ".js-notice-block__title", ".price_value"),
    "dns-shop.ru": Rule(".catalog-product", ".catalog-product__name", ".product-buy__price", rating=".catalog-product__rating b", delivery=".delivery-info-widget__button"),
    "etm.ru": Rule('[class*="listItem-item"]', '[data-testid="link-product-name"]', '[data-testid^="catalog-list-item-price-details-"]:not([data-testid$="-unit"])', delivery='[class*="item-breakSpaces"]'),
    "komus.ru": Rule(".product-card-list, .product-card, .product-card-list__description", ".product-title", ".v-product-price__value", rating=".product-rating", delivery=".scenario-availability-item__in-stock"),
    "labirint.ru": Rule(".product-card[data-product-id]", ".product-card__name", ".product-card__price-current"),
    "lemanapro.ru": Rule('[data-qa="product"], [data-qa="product-card"], [data-qa="products-list"]', '[data-qa="product-name"]', '[data-testid="price-integer"]'),
    "market.yandex.ru": Rule('[data-apiary-widget-name="@marketfront/SerpEntity"], [data-zone-name="productSnippet"]', '[data-auto="snippet-title"]', '[data-auto="snippet-price-current"], [class*="ds-text_headline-5"]', 'a[href*="/card/"], a[href*="/product--/"]'),
    "mvideo.ru": Rule("a[mvid-product-card]", ".name", ".current-price", ":self"),
    "oaopolimer.ru": Rule(".product-item-big-card", ".title a", ".price"),
    "officemag.ru": Rule(".js-productListItem, .BannerProduct", ".listItemTitle a, .listItem__title a, a[href*='/catalog/goods/']", ".ProductSpecial .Price, .listItemBuy__price"),
    "ozon.ru": Rule(".tile-root", ".tsBody500Medium, .tsBody400Small", ".tsHeadline500Medium", 'a[href*="/product/"]'),
    "petrovich.ru": Rule('[data-test="product-card-catalog-wide"], [data-test="product-card-catalog"]', '[data-test="product-title"]', '[data-test="product-price"], [data-test="price"]', '[data-test="product-link"]'),
    "rs24.ru": Rule(".search-results__item.js-product", ".item-description a[title], a[href^='/product/'][title]", ".js-product-price", 'a[href^="/product/"]'),
    "tara.ru": Rule(".catalog-block-view__item", ".item-title a, .js-notice-block__title", '.price_value, [itemprop="price"]'),
    "spb.tara.ru": Rule(".catalog-block-view__item", ".item-title a, .js-notice-block__title", '.price_value, [itemprop="price"]'),
    "spb.kuvalda.ru": Rule(".catalog__list-item.snippet, .promo-slider__item.snippet", ".snippet__title", ".snippet-price__value"),
    "vseinstrumenti.ru": Rule('[data-qa="products-tile"]', '[data-qa="product-name"]', '[data-qa="product-price-current"]', 'a[data-qa="product-name"], a[data-qa="product-photo-click"]', rating='[data-qa="product-rating"]'),
    "wildberries.ru": Rule("article.product-card", ".product-card__name", ".price__lower-price", ".product-card__link", rating=".address-rate-mini", delivery='[data-helper="delivery-display"]'),
    "ximtek.ru": Rule(".product-item-small-card", ".product-item-title a", ".product-item-price-current"),
    "xn----7sbbnaebi2cxajxwo0b.xn--p1ai": Rule(".product-list-item", ".product-list-item-title", ".product-list-item-price"),
}
RULES["спецодежда-тула.рф"] = RULES["xn----7sbbnaebi2cxajxwo0b.xn--p1ai"]


@dataclass
class Candidate:
    values: dict
    method: str
    evidence: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def candidate(name, price, url, base, method, **extra) -> Candidate | None:
    name = clean(name)
    url = canonical_url(url, base)
    if not name or not url:
        return None
    values = {"is_product": 1, "name_gt": name, "price_gt": number(price), "product_url_gt": url}
    values.update({k: v for k, v in extra.items() if v is not None})
    return Candidate(values, method)


def lemanapro_state(soup: BeautifulSoup, base: str) -> list[Candidate]:
    result = []
    pattern = re.compile(r'window\.INITIAL_STATE\["plp"\]\s*=\s*')
    for script in soup.find_all("script"):
        text = script.get_text()
        match = pattern.search(text)
        if not match:
            continue
        state, _ = json.JSONDecoder().raw_decode(text, match.end())
        products = state.get("plp", {}).get("plp", {}).get("products") or {}
        ids = set(map(str, products.get("productsIds") or []))
        # Only the listing's product array, never navigation/recommendation data.
        for row in products.get("productsData") or []:
            if ids and str(row.get("productId")) not in ids:
                continue
            pricing = row.get("price") or {}
            value = candidate(row.get("displayedName"), pricing.get("main_price"), row.get("productLink"), base, "js:INITIAL_STATE.plp.plp.plp.products.productsData")
            if value:
                value.evidence = {"productId": row.get("productId"), "price": pricing}
                result.append(value)
    return result


def extract_cards(html_text: str, domain: str, base: str) -> tuple[list[Candidate], dict]:
    soup = BeautifulSoup(html_text, "lxml")
    rule = RULES.get(domain)
    if not rule:
        return [], {"reason": "unsupported_domain", "card_nodes": 0}
    cards = soup.select(rule.cards)
    if domain == "lemanapro.ru":
        cards = soup.select('[data-qa="product"], [data-qa="product-card"]') or cards
    # Avoid double extraction of nested card/description containers.
    selected_ids = {id(card) for card in cards}
    cards = [card for card in cards if not any(id(p) in selected_ids for p in card.parents)]
    results = []
    for index, card in enumerate(cards, 1):
        name_node = first(card, rule.name)
        name = (clean(name_node.get("title")) or text_of(name_node)) if name_node is not None else ""
        link_node = first(card, rule.link) if rule.link else name_node
        if link_node is not None and not link_node.get("href"):
            link_node = link_node.find("a", href=True) or link_node.find_parent("a", href=True)
        url = link_node.get("href") if link_node is not None else None
        price_node = first(card, rule.price)
        price_text = text_of(price_node)
        extra = {}
        warnings = []
        if domain == "rs24.ru":
            name = clean(card.get("data-description")) or name
            price_text = card.get("data-retail-price") or price_text
        elif domain == "officemag.ru":
            # The analytics item explicitly gives the single-unit price. Display
            # price may instead be the lowest volume tier (e.g. 342 vs 349 RUB).
            try:
                item = json.loads(card.get("data-ga-object", "{}"))["items"][0]
                name = BeautifulSoup(item["name"], "lxml").get_text("", strip=True)
                price_text = item.get("price")
            except (ValueError, KeyError, IndexError, TypeError):
                pass
            if "BannerProduct" in card.get("class", []):
                image = card.select_one("img[alt]")
                name = clean(image.get("alt")) if image else name
                url = card.get("href")
                price_text = text_of(card.select_one(".Price__count")) + "." + (text_of(card.select_one(".Price__penny")) or "00")
                warnings.append("promotional_banner_product")
            else:
                warnings.append("price_for_one_unit; check quantity tiers")
        elif domain == "wildberries.ru":
            name = clean(link_node.get("aria-label")) if link_node else name
            if card.select_one(".wallet-price"):
                warnings.append("price_requires_WB_wallet")
        elif domain == "market.yandex.ru":
            if "Пэй" in card.get_text():
                warnings.append("price_may_require_Yandex_Pay")
        elif domain == "petrovich.ru" and not price_text:
            for node in card.find_all(attrs={"data-test": re.compile("price")}):
                if "₽" in node.get_text() and number(node.get_text()) is not None:
                    price_text = text_of(node)
                    break
        elif domain == "lemanapro.ru":
            # Old/current prices use the same inner price-integer selector.
            for node in card.select('[data-testid="price-integer"]'):
                if node.find_parent(attrs={"data-testid": "price-block-oldprice"}) is None:
                    price_text = text_of(node)
                    break
        elif domain == "komus.ru":
            minimum = re.search(r"от\s+\d+\s+шт\.", card.get_text(" ", strip=True))
            if minimum:
                warnings.append("minimum_order: " + clean(minimum.group()))
        elif domain in {"xn----7sbbnaebi2cxajxwo0b.xn--p1ai", "спецодежда-тула.рф"}:
            price_text = next((text_of(n) for n in card.select(".product-list-item-price") if "Цена:" in n.get_text()), "")
        if rule.rating:
            rating = number(text_of(first(card, rule.rating)))
            if rating is not None and 0 < rating <= 5:
                extra["rating_gt"] = rating
        if rule.delivery:
            delivery = text_of(first(card, rule.delivery))
            if delivery:
                extra["delivery_time_gt"] = delivery
        stock_text = clean(card.get_text(" ", strip=True)).casefold()
        availability = card.select_one('[itemprop="availability"]')
        availability_value = clean(availability.get("href") or availability.get("content")) if availability is not None else ""
        if "OutOfStock" in availability_value or "Discontinued" in availability_value:
            extra["in_stock_gt"] = 0
        elif "InStock" in availability_value:
            extra["in_stock_gt"] = 1
        elif re.search(r"нет в наличии|нет в продаже|товар закончился|снят с продажи", stock_text):
            extra["in_stock_gt"] = 0
        elif re.search(r"\bв наличии\b|доступно сегодня|есть в наличии", stock_text):
            extra["in_stock_gt"] = 1
        if number(price_text) == 0 and re.search(r"цен[аыу].{0,35}по запросу|запросите цену", stock_text):
            price_text = None
            warnings.append("price_on_request; zero in metadata is not a price")
        value = candidate(name, price_text, url, base, f"css:{rule.cards} [{index}]", **extra)
        if value:
            value.evidence = {"name_selector": rule.name, "price_selector": rule.price, "price_text": price_text}
            value.warnings.extend(warnings)
            results.append(value)
    if domain == "lemanapro.ru":
        # JSON holds the full list even when only the first cards are rendered.
        results.extend(lemanapro_state(soup, base))
    by_url = {}
    conflicts = []
    for value in results:
        url = value.values["product_url_gt"]
        if url not in by_url:
            by_url[url] = value
            continue
        existing = by_url[url]
        for key, val in value.values.items():
            if blank(existing.values.get(key)):
                existing.values[key] = val
            elif key in {"price_gt", "name_gt"} and not blank(val) and existing.values[key] != val:
                conflicts.append({"url": url, "field": key, "kept": existing.values[key], "alternative": val})
                existing.warnings.append(f"conflicting_{key}; check source")
    rejected = len(cards) - sum(1 for c in results if c.method.startswith("css:"))
    return list(by_url.values()), {"card_nodes": len(cards), "rejected_cards": rejected, "conflicts": conflicts, "reason": "extracted" if by_url else "no_cards_extracted"}


def records(sheet) -> list[dict]:
    rows = iter(sheet.values)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


def extract_workbook(workbook, directory: Path) -> tuple[dict, dict]:
    extracted = {}
    pages_report = []
    pages = records(workbook["Pages"])
    for index, page in enumerate(pages, 1):
        relpath = page["benchmark_html_file"]
        path = (directory / relpath).resolve()
        if not path.is_relative_to(directory.resolve()):
            raise ValueError(f"HTML path outside dataset: {relpath}")
        domain = page["source"]
        report = {"html_file": relpath, "page_id": page["page_id"], "source": domain}
        if not path.is_file():
            candidates, diagnostics = [], {"reason": "missing_html", "card_nodes": 0}
        else:
            raw = path.read_bytes()
            candidates, diagnostics = extract_cards(raw.decode("utf-8", errors="replace"), domain, page.get("page_url") or f"https://{domain}/")
            report["html_sha256"] = hashlib.sha256(raw).hexdigest()
        extracted[relpath] = candidates
        report.update(diagnostics)
        report["products"] = len(candidates)
        report["missing_price"] = sum(c.values.get("price_gt") is None for c in candidates)
        report["candidates"] = [{"values": c.values, "method": c.method, "evidence": c.evidence, "warnings": c.warnings} for c in candidates]
        pages_report.append(report)
        if index % 20 == 0 or index == len(pages):
            print(f"Extracted {index}/{len(pages)} pages", flush=True)
    domain_totals = defaultdict(Counter)
    for page in pages_report:
        counts = domain_totals[page["source"]]
        counts.update(pages=1, products=page["products"], pages_with_products=int(page["products"] > 0), missing_price=page["missing_price"])
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "status": "machine_prefill_requires_manual_review", "pages": pages_report,
        "by_source": {key: dict(value) for key, value in domain_totals.items()},
        "total_products": sum(p["products"] for p in pages_report),
        "pages_with_products": sum(p["products"] > 0 for p in pages_report),
    }
    return extracted, report


def fill_workbook(workbook, extracted: dict, report: dict) -> dict:
    sheet = workbook["GroundTruth"]
    headers = [cell.value for cell in sheet[1]]
    originals = records(sheet)
    source_rows = list(sheet.iter_rows(min_row=2))
    groups = defaultdict(list)
    for index, row in enumerate(originals):
        groups[row["html_file"]].append((index, row))
    next_id = max((r["gt_id"] for r in originals if isinstance(r["gt_id"], int)), default=0) + 1
    planned = []
    added, filled = 0, 0
    for html_file, group in groups.items():
        rows = [(index, dict(row), {}) for index, row in group]
        # Reviewed or explicitly negative pages must never gain inferred labels.
        protected = any(row.get("annotation_status") in {"done", "skip"} or row.get("is_product") == 0 for _, row in group)
        if not protected:
            for value in extracted.get(html_file, []):
                url = value.values["product_url_gt"]
                match = next((entry for entry in rows if canonical_url(entry[1].get("product_url_gt"), url) == url), None)
                if match is None:
                    match = next((entry for entry in rows if all(blank(entry[1].get(k)) for k in PRODUCT_FIELDS) and blank(entry[1].get("annotator"))), None)
                if match is None:
                    template_index, template = group[0]
                    row = {key: template.get(key) for key in ("page_id", "html_file", "source", "run_id", "query", "page_url", "page_kind")}
                    row["gt_id"] = next_id
                    next_id += 1
                    match = (template_index, row, {})
                    rows.append(match)
                    added += 1
                _, row, comments = match
                changes = []
                for key, val in value.values.items():
                    if not blank(val) and blank(row.get(key)):
                        row[key] = val
                        changes.append(key)
                        comments[key] = f"Automatic candidate; verify manually.\nHTML: {html_file}\nMethod: {value.method}\nEvidence: {json.dumps(value.evidence, ensure_ascii=False)}"
                        filled += 1
                if changes:
                    if blank(row.get("entity_id")):
                        row["entity_id"] = "auto-" + hashlib.sha256((html_file + "\n" + url).encode()).hexdigest()[:16]
                    if blank(row.get("annotator")):
                        row["annotator"] = ANNOTATOR
                    if row.get("annotation_status") in {None, "", "pending"}:
                        row["annotation_status"] = "in_progress"
                    note = "AUTO PREFILL — requires manual review; " + value.method
                    if value.warnings:
                        note += "; " + "; ".join(sorted(set(value.warnings)))
                    if note not in (row.get("notes") or ""):
                        row["notes"] = "; ".join(filter(None, [row.get("notes"), note]))
        planned.extend(rows)
    if not filled:
        return {"added_rows": 0, "filled_cells": 0, "groundtruth_rows": len(originals)}
    # Snapshot cell properties before expanding/reordering the sheet.
    templates = [[(copy(c._style), copy(c.comment), copy(c.hyperlink), c.value, c.data_type) for c in row] for row in source_rows]
    for output_row, (template_index, values, comments) in enumerate(planned, 2):
        for col, key in enumerate(headers, 1):
            cell = sheet.cell(output_row, col)
            style, comment, hyperlink, old_value, old_type = templates[template_index][col - 1]
            cell._style = copy(style)
            cell.value = values.get(key)
            # HTML strings are text, never Excel formulas.
            if cell.value == old_value:
                cell.data_type = old_type
            elif isinstance(cell.value, str):
                cell.data_type = "s"
            cell.comment = Comment(comments[key], ANNOTATOR) if key in comments else copy(comment)
            cell.hyperlink = copy(hyperlink)
            if key == "price_gt":
                cell.number_format = "0.00"
    last_col = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_col}{sheet.max_row}"
    for validation in sheet.data_validations.dataValidation:
        for cell_range in validation.sqref.ranges:
            cell_range.max_row = max(cell_range.max_row, sheet.max_row)
    for table in sheet.tables.values():
        table.ref = sheet.auto_filter.ref
    if "GT_Prefill_Audit" in workbook:
        del workbook["GT_Prefill_Audit"]
    audit = workbook.create_sheet("GT_Prefill_Audit")
    audit.append(["html_file", "source", "status", "card_nodes", "products", "missing_price", "rejected_cards", "conflicts"])
    for page in report["pages"]:
        audit.append([page.get(k) for k in ("html_file", "source", "reason", "card_nodes", "products", "missing_price", "rejected_cards")] + [len(page.get("conflicts", []))])
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    for cell in audit[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for column, width in {"A": 62, "B": 30, "C": 25, "D": 15, "E": 15, "F": 18, "G": 18, "H": 15}.items():
        audit.column_dimensions[column].width = width
    return {"added_rows": added, "filled_cells": filled, "groundtruth_rows": len(planned)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, default=Path(__file__).resolve().parents[1] / "industrial" / "dind.xlsx")
    parser.add_argument("--write", action="store_true", help="Back up Excel and apply missing fields")
    parser.add_argument("--report", type=Path, help="JSON diagnostics including every extracted candidate")
    parser.add_argument("--metadata-only", action="store_true", help="Recover query without re-extracting or adding products")
    parser.add_argument("--confirm-empty-pending", action="store_true", help="Apply human confirmation that blank pending rows represent pages without products")
    parser.add_argument("--confirm-unpriced-unavailable", action="store_true", help="Apply human confirmation that products with blank/zero prices are no longer sold; preserve price cells")
    args = parser.parse_args()
    path = args.workbook.resolve()
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path)
    if args.metadata_only:
        rows = records(workbook["GroundTruth"])
        positive = [r for r in rows if r.get("is_product") == 1]
        report = {"total_products": len(positive), "pages_with_products": len({r["html_file"] for r in positive}),
                  "by_source": {}, "status": "metadata_update", "generated_at": datetime.now(timezone.utc).isoformat(),
                  "changes": {"added_rows": 0, "filled_cells": 0, "groundtruth_rows": len(rows)}}
    else:
        extracted, report = extract_workbook(workbook, path.parent)
        report["changes"] = fill_workbook(workbook, extracted, report)
    report["workbook"] = str(path)
    report["workbook_sha256_before"] = before_hash
    report["metadata"] = update_metadata(workbook, path.parent,
        confirm_empty_pending=args.confirm_empty_pending,
        confirm_unpriced_unavailable=args.confirm_unpriced_unavailable)
    report["written"] = False
    report_path = args.report or path.with_name("gt_metadata_report.json" if args.metadata_only else "gt_prefill_report.json")
    changed = report["changes"]["filled_cells"] or any(report["metadata"][key] for key in (
        "query_cells_filled", "confirmed_empty_rows", "confirmed_unavailable_rows", "audit_cells_updated"))
    if args.write and changed:
        if hashlib.sha256(path.read_bytes()).hexdigest() != before_hash:
            raise RuntimeError("Workbook changed during extraction; retry against the latest file")
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        backup = path.with_name(f"{path.stem}.before-prefill-{stamp}{path.suffix}")
        shutil.copy2(path, backup)
        temp = path.with_name(f".{path.stem}.prefill-{stamp}.xlsx")
        workbook.save(temp)
        check = load_workbook(temp, read_only=True)
        assert check["GroundTruth"].max_row == report["changes"]["groundtruth_rows"] + 1
        check.close()
        report["backup"] = str(backup)
        try:
            temp.replace(path)
        except PermissionError as exc:
            report["prepared_workbook"] = str(temp)
            report["write_error"] = str(exc)
            report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            raise SystemExit(
                f"Cannot replace {path}; close it in Excel or your editor. "
                f"Prepared workbook: {temp}. Report: {report_path}"
            ) from exc
        report["written"] = True
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: report[k] for k in ("total_products", "pages_with_products", "changes", "written")}, ensure_ascii=False, indent=2))
    print("Metadata:", json.dumps({k: v for k, v in report["metadata"].items() if k != "pages"}))
    for domain, totals in report["by_source"].items():
        print(domain, dict(totals))
    print(f"Report: {report_path}")
    if report.get("backup"):
        print(f"Backup: {report['backup']}")


if __name__ == "__main__":
    main()
