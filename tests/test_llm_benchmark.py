"""Offline tests of the actual notebook. No key or paid API requests are needed."""
import contextlib
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.run_llm_benchmark import load_notebook


def product(**changes):
    result = dict(name='Дрель Alpha 500', price=100, old_price=None, product_url='/product/1',
                  image_url=None, availability=True, delivery_time=None, reviews_count=None)
    return result | changes


def response(products=None, *, content=None, finish='stop', status=200, usage=True):
    body = dict(id='fake-response', model='qwen/qwen3.8-omni-flash',
                choices=[dict(finish_reason=finish, message=dict(content=content if content is not None else json.dumps({'products': products or []}, ensure_ascii=False)))])
    if usage:
        body['usage'] = dict(prompt_tokens=100, completion_tokens=20, completion_tokens_details={'reasoning_tokens': 0})
    if status != 200:
        body = {'error': {'message': 'simulated provider error'}}
    return dict(body=json.dumps(body), http_status=status, headers={}, transport_error=None, latency_ms=12)


class LLMBaselineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with contextlib.redirect_stdout(io.StringIO()):
            cls.ns = load_notebook()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.settings = patch.dict(self.ns, BENCHMARK_ROOT=self.root, OUTPUT_DIR=self.root / 'out',
                                   RUN_MODE='run', RETRY_ERRORS=False, RETRY_UNCERTAIN=False,
                                   PAGE_LIMIT=None, RUN_DATASETS=('synthetic',), EXPERIMENT_ID='test',
                                   LLM_CONFIG=self.ns['LLM_CONFIG'] | {'max_attempts': 2})
        self.settings.start()
        self.addCleanup(self.settings.stop)

    def page(self):
        path = self.root / 'page.html'
        path.write_text('<html><script>window.state={"product":"Дрель"}</script>FULL_TAIL</html>', encoding='utf-8')
        return dict(dataset='synthetic', page_id='p1', html_file='html/page.html', html_key='page.html',
                    html_path=str(path), source='demo.test', run_id='1', skip_reason='', query='irrelevant search',
                    page_url='https://demo.test/search', source_url='https://demo.test',
                    content_sha256=hashlib.sha256(path.read_bytes()).hexdigest())

    def test_full_html_no_gt_and_reusable_response(self):
        page = self.page()
        calls = []
        def fake(payload):
            calls.append(payload)
            return response([product()])
        result = self.ns['llm_page_call'](page, fake)
        self.assertEqual(result['rows'][0]['product_url'], 'https://demo.test/product/1')
        self.assertTrue(calls[0]['messages'][1]['content'].endswith(Path(page['html_path']).read_text(encoding='utf-8')))
        self.assertEqual(calls[0]['model'], 'qwen/qwen3.8-omni-flash')
        self.assertEqual(calls[0]['reasoning'], {'effort': 'none', 'exclude': True})
        cached = self.ns['llm_page_call'](page, lambda _: self.fail('cached page called API'))
        self.assertTrue(cached['reused'])
        self.assertEqual(cached['new_attempts'], 0)
        page['query'] = 'changed'
        with self.assertRaisesRegex(RuntimeError, 'differs'):
            self.ns['llm_page_call'](page, fake)

    def test_recovers_response_before_checkpoint_and_detects_tampering(self):
        page = self.page()
        self.ns['llm_page_call'](page, lambda _: response([product()]))
        folder = next((self.root / 'out/synthetic/llm_calls').iterdir())
        (folder / 'result.json').unlink()
        result = self.ns['llm_page_call'](page, lambda _: self.fail('recovery called API'))
        self.assertTrue(result['reused'])
        next(folder.glob('*.response.json.gz')).write_bytes(b'corrupt')
        with self.assertRaisesRegex(RuntimeError, 'modified'):
            self.ns['llm_page_call'](page, lambda _: self.fail('tamper called API'))

    def test_uncertain_attempt_needs_explicit_retry(self):
        page = self.page()
        with self.assertRaises(KeyboardInterrupt):
            self.ns['llm_page_call'](page, lambda _: (_ for _ in ()).throw(KeyboardInterrupt()))
        with self.assertRaisesRegex(RuntimeError, 'unknown billing'):
            self.ns['llm_page_call'](page, lambda _: self.fail('uncertain retry'))
        with patch.dict(self.ns, RETRY_UNCERTAIN=True):
            result = self.ns['llm_page_call'](page, lambda _: response([]))
        self.assertEqual(len(result['attempts']), 2)
        self.assertEqual(result['attempts'][0]['status'], 'interrupted_unknown')

    def test_retry_accounting_and_unknown_usage(self):
        replies = iter([response(status=429), response([product()], usage=False)])
        with patch('time.sleep'):
            record = self.ns['llm_page_call'](self.page(), lambda _: next(replies))
        self.assertEqual(record['new_attempts'], 2)
        cost = self.ns['llm_cost_summary']([record])
        self.assertEqual(cost['http_attempts'], 2)
        self.assertIsNone(cost['prompt_tokens']['known_total'])
        self.assertEqual(cost['prompt_tokens']['unknown_attempts'], 2)

    def test_invalid_truncated_empty_and_optional_schema(self):
        page = self.page()
        decode = self.ns['llm_decode_response']
        self.assertEqual(decode(response(content='{bad'), page)['status'], 'invalid_json_or_schema')
        self.assertEqual(decode(response([product()], finish='length'), page)['status'], 'output_truncated')
        self.assertEqual(decode(response([]), page)['status'], 'ok')
        self.assertEqual(decode(response([product(price=None)]), page)['rows'][0]['pred_price'], None)
        invalid = decode(response([product(), product(product_url='javascript:alert(1)')]), page)
        self.assertEqual(invalid['status'], 'partial_validation')
        self.assertEqual(len(invalid['rows']), 1)
        self.assertEqual(len(invalid['rejected']), 1)

    def test_403_stops_before_next_page(self):
        page = self.page()
        manifest = pd.DataFrame([page, page | {'page_id': 'p2'}])
        calls = []
        with contextlib.redirect_stdout(io.StringIO()), self.assertRaisesRegex(RuntimeError, 'access'):
            self.ns['run_llm_baseline'](manifest, lambda payload: calls.append(payload) or response(status=403))
        self.assertEqual(len(calls), 1)

    def test_key_loading_is_local_and_does_not_modify_environment(self):
        with patch.dict('os.environ', {}, clear=True):
            (self.root / '.env').write_text('OTHER=ignore\nROUTERAI_API_KEY="unit-test-not-a-real-key" # comment\n')
            self.assertEqual(self.ns['llm_api_key'](), 'unit-test-not-a-real-key')
            import os
            self.assertNotIn('ROUTERAI_API_KEY', os.environ)

    def test_failed_empty_page_is_not_counted_as_success(self):
        evaluation = pd.DataFrame(columns=self.ns['EXPORT_EVALUATION_COLUMNS'])
        pages = pd.DataFrame([dict(dataset='demo', variant='llm_full_html', page_id='empty',
                                   is_scored=True, gt_positive_count=0, extraction_status='invalid_json_or_schema')])
        _, _, metrics = self.ns['paper_evaluate'](evaluation, pages)
        self.assertEqual(metrics.PSR_empty.iloc[0], 0)
        self.assertEqual(metrics.page_exact_rate.iloc[0], 0)

    def test_strict_metrics_keep_price_absence_and_do_not_reuse_tolerant_price(self):
        page = self.page()
        manifest = pd.DataFrame([page])
        reference = self.root / 'gt.csv'
        pd.DataFrame([
            dict(page_id='p1', html_file='html/page.html', entity_id='e1', is_product=1,
                 name_gt='Дрель Alpha 500', price_gt=100, product_url_gt='https://demo.test/product/1'),
            dict(page_id='p1', html_file='html/page.html', entity_id='e2', is_product=1,
                 name_gt='Unavailable drill', price_gt=0, product_url_gt='https://demo.test/product/2'),
        ]).to_csv(reference, index=False)
        gt = self.ns['read_groundtruth'](reference, zero_price_is_missing=True)
        with contextlib.redirect_stdout(io.StringIO()):
            raw = self.ns['run_llm_baseline'](manifest, lambda _: response([
                product(price=100.5), product(name='Unavailable drill', price=None, product_url='/product/2')]))
        evaluation, ordinary = self.ns['evaluate_predictions'](raw, gt, manifest)
        pages = self.ns['build_page_analytics'](evaluation, gt, manifest, raw)
        entities, _, strict = self.ns['paper_evaluate'](evaluation, pages)
        self.assertEqual(ordinary.price_TP.iloc[0], 1)
        self.assertEqual(strict.price_TP.iloc[0], 0)
        self.assertEqual(strict.price_FN.iloc[0], 1)
        self.assertEqual(entities.strict_acceptable_entity.sum(), 1)

    def test_shared_evaluation_cells_are_identical_to_original(self):
        original = json.loads((ROOT / 'baseline_determined.ipynb').read_text(encoding='utf-8'))
        notebook = json.loads((ROOT / 'baseline_llm.ipynb').read_text(encoding='utf-8'))
        code = [''.join(c['source']) for c in notebook['cells'] if c['cell_type'] == 'code']
        for index in [6, 8, 14]:
            self.assertIn(''.join(original['cells'][index]['source']), code)

    def test_complete_run_and_offline_replay_export_excel(self):
        folder = self.root / 'synthetic'
        (folder / 'html').mkdir(parents=True)
        for name, text in [('a', 'PRODUCT'), ('b', 'EMPTY')]:
            (folder / 'html' / (name + '.html')).write_text(text)
        reference = folder / 'synthetic_ground_truth.csv'
        pd.DataFrame([
            dict(page_id='a', entity_id='e1', html_file='html/a.html', is_product=1,
                 name_gt='Дрель Alpha 500', price_gt=100, product_url_gt='https://demo.test/product/1', in_stock_gt=True),
            dict(page_id='b', entity_id='', html_file='html/b.html', is_product=0, name_gt='', price_gt=None, product_url_gt='')
        ]).to_csv(reference, index=False)
        pd.DataFrame([dict(page_id=name, html_file=f'html/{name}.html', source='demo.test', page_url='https://demo.test/search')
                      for name in ['a', 'b']]).to_csv(folder / 'synthetic_pages.csv', index=False)
        calls = []
        def fake(payload):
            calls.append(payload)
            return response([product()] if payload['messages'][1]['content'].endswith('PRODUCT') else [])
        with patch.dict(self.ns, DATASETS={'synthetic': dict(root=folder, html_dir=folder/'html', reference_file=reference)},
                        llm_prepare_experiment=lambda: None, llm_transport=fake), contextlib.redirect_stdout(io.StringIO()):
            result = self.ns['llm_main']()
            self.assertEqual(result['status'], 'complete')
            self.ns['RUN_MODE'] = 'replay'
            self.ns['llm_main']()
        self.assertEqual(len(calls), 2)
        output = self.root / 'baseline_results/llm/test'
        metrics = pd.read_csv(output / 'synthetic/baseline_metrics.csv')
        self.assertEqual(metrics.name_F1.iloc[0], 1)
        strict = pd.read_csv(output / 'synthetic/paper_required_metrics.csv')
        self.assertEqual(strict.PSR_empty.iloc[0], 1)
        self.assertEqual(strict.PSR_positive.iloc[0], 1)
        workbook = load_workbook(output / 'baseline_comparison.xlsx', read_only=True)
        self.assertIn('Сравнение', workbook.sheetnames)
        self.assertEqual(workbook['Сравнение'].max_row, 2)
        workbook.close()
        cost = json.loads((output / 'synthetic/llm_cost_summary.json').read_text())
        self.assertEqual(cost['new_http_attempts'], 0)
        self.assertEqual(cost['http_attempts'], 2)


if __name__ == '__main__':
    unittest.main()
