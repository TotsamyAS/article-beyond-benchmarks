"""Offline checks: python -m unittest tests.test_prefill_industrial_gt -v."""
import json
from pathlib import Path
import unittest

from openpyxl import Workbook

from scripts.prefill_industrial_gt import (
    ANNOTATOR, Candidate, extract_cards, fill_workbook, number, records,
    query_from_url, recover_query, update_metadata,
)
from tempfile import TemporaryDirectory
from urllib.parse import quote


class CardExamplesTest(unittest.TestCase):
    def test_supplied_examples(self):
        expected = {
            "bigam.ru": 4170, "chipdip.ru": 2490, "citilink.ru": 11600,
            "cnc.su": 2867.14, "dns-shop.ru": 7199, "etm.ru": 3279,
            "komus.ru": None, "labirint.ru": 765, "lemanapro.ru": 7490,
            "market.yandex.ru": 2868, "mvideo.ru": 56499, "oaopolimer.ru": 1000,
            "officemag.ru": 349, "ozon.ru": 7599, "petrovich.ru": 6190,
            "rs24.ru": 223.26, "vseinstrumenti.ru": 6621,
            "wildberries.ru": 1403, "ximtek.ru": 2313.61,
        }
        root = Path(__file__).resolve().parents[1] / "industrial/html/cadr_examples"
        if not root.is_dir():
            self.skipTest("Historical card_examples fixtures were removed after GT review")
        for domain, price in expected.items():
            with self.subTest(domain=domain):
                cards, info = extract_cards((root / f"{domain}.html").read_text(encoding="utf-8"), domain, f"https://{domain}/")
                self.assertEqual(len(cards), 2 if domain == "petrovich.ru" else 1)
                self.assertEqual(cards[0].values["price_gt"], price)
                self.assertTrue(cards[0].values["name_gt"])
                self.assertTrue(cards[0].values["product_url_gt"].startswith("https://"))
                self.assertEqual(info["rejected_cards"], 0)

    def test_js_listing_excludes_other_arrays_and_handles_quoted_braces(self):
        item = {"productId": "1", "displayedName": 'Tool "}];"', "productLink": "/product/tool-1/", "price": {"main_price": 10, "previous_price": 15}}
        state = {"plp": {"plp": {"products": {"productsIds": ["1"], "productsData": [item]}, "recommendations": [dict(item, productId="2")]}}}
        source = '<script>window.INITIAL_STATE["plp"]=' + json.dumps(state) + ';alert("must not run");</script>'
        cards, _ = extract_cards(source, "lemanapro.ru", "https://lemanapro.ru/")
        self.assertEqual(len(cards), 1)
        self.assertEqual(cards[0].values["name_gt"], item["displayedName"])
        self.assertEqual(cards[0].values["price_gt"], 10)

    def test_no_card_is_not_negative_annotation(self):
        cards, info = extract_cards("<h1>No results</h1>", "dns-shop.ru", "https://dns-shop.ru/")
        self.assertEqual(cards, [])
        self.assertEqual(info["reason"], "no_cards_extracted")

    def test_observed_alternate_layouts(self):
        root = Path(__file__).resolve().parents[1] / "industrial/html"
        for filename, count, price in [
            ("bigam.ru-run-5ce6eea7.html", 15, 100),
            ("lemanapro.ru-run-none3.html", 60, 3410),
            ("officemag.ru-run-none2.html", 4, 1715.95),
            ("tara.ru-run-08e7e74c.html", 20, None),
        ]:
            with self.subTest(filename=filename):
                domain = filename.split("-run-")[0]
                cards, info = extract_cards((root / filename).read_text(encoding="utf-8"), domain, f"https://{domain}/")
                self.assertEqual(len(cards), count)
                self.assertEqual(cards[0].values["price_gt"], price)
                self.assertFalse(info["conflicts"])

    def test_bigam_regular_price_without_discount(self):
        source = '<div class="a-product-card"><h3 class="a-product-card__title"><a href="/product/tool/">Tool</a></h3><div class="a-price__current">48 590</div></div>'
        cards, _ = extract_cards(source, "bigam.ru", "https://bigam.ru/")
        self.assertEqual(cards[0].values["price_gt"], 48590)

    def test_empty_js_product_list(self):
        source = '<script>window.INITIAL_STATE["plp"]={"plp":{"plp":{"products":{"productsData":null}}}};</script>'
        cards, _ = extract_cards(source, "lemanapro.ru", "https://lemanapro.ru/")
        self.assertEqual(cards, [])

    def test_prices_do_not_concatenate_old_price_or_currency(self):
        for raw, expected in [("2\u00a0313.61 ₽", 2313.61), ("349 , 00 руб.", 349), ("1000 руб. 00 коп.", 1000), ("7 199 ₽ 9 999 ₽", 7199), ("Цена по запросу", None), (0, 0)]:
            with self.subTest(raw=raw):
                self.assertEqual(number(raw), expected)


class WorkbookTest(unittest.TestCase):
    def setUp(self):
        self.book = Workbook()
        self.sheet = self.book.active
        self.sheet.title = "GroundTruth"
        self.headers = ["gt_id", "page_id", "html_file", "entity_id", "source", "run_id", "query", "page_url", "page_kind", "is_product", "name_gt", "price_gt", "product_url_gt", "in_stock_gt", "rating_gt", "delivery_time_gt", "region_gt", "annotator", "annotation_status", "notes"]
        self.sheet.append(self.headers)
        self.add_row(gt_id=1, html_file="html/demo.html", page_id="page-1", source="demo", annotation_status="pending")
        self.book.create_sheet("Pages").append(["untouched"])
        self.candidates = [Candidate({"is_product": 1, "name_gt": "Tool", "price_gt": 10, "product_url_gt": "https://demo/p/1"}, "test"), Candidate({"is_product": 1, "name_gt": "Tool 2", "price_gt": 20, "product_url_gt": "https://demo/p/2"}, "test")]
        self.report = {"pages": []}

    def add_row(self, **values):
        self.sheet.append([values.get(k) for k in self.headers])

    def test_expansion_preserves_ids_and_is_idempotent(self):
        report = fill_workbook(self.book, {"html/demo.html": self.candidates}, self.report)
        self.assertEqual(report["added_rows"], 1)
        rows = records(self.sheet)
        self.assertEqual([r["gt_id"] for r in rows], [1, 2])
        self.assertEqual([r["annotator"] for r in rows], [ANNOTATOR, ANNOTATOR])
        self.assertEqual([r["annotation_status"] for r in rows], ["in_progress"] * 2)
        repeated = fill_workbook(self.book, {"html/demo.html": self.candidates}, self.report)
        self.assertEqual(repeated["filled_cells"], 0)
        self.assertEqual(records(self.sheet), rows)
        self.assertEqual(self.book["Pages"]["A1"].value, "untouched")

    def test_existing_label_is_preserved_while_missing_price_is_filled(self):
        self.sheet.cell(2, self.headers.index("name_gt") + 1, "Human name")
        self.sheet.cell(2, self.headers.index("product_url_gt") + 1, "https://demo/p/1")
        self.sheet.cell(2, self.headers.index("annotator") + 1, "Human")
        fill_workbook(self.book, {"html/demo.html": self.candidates[:1]}, self.report)
        row = records(self.sheet)[0]
        self.assertEqual(row["name_gt"], "Human name")
        self.assertEqual(row["annotator"], "Human")
        self.assertEqual(row["price_gt"], 10)

    def test_done_and_negative_pages_are_not_expanded(self):
        for status, is_product in [("done", None), ("pending", 0), ("skip", None)]:
            with self.subTest(status=status, is_product=is_product):
                self.sheet.cell(2, self.headers.index("annotation_status") + 1).value = status
                self.sheet.cell(2, self.headers.index("is_product") + 1).value = is_product
                before = records(self.sheet)
                result = fill_workbook(self.book, {"html/demo.html": self.candidates}, self.report)
                self.assertEqual(result["filled_cells"], 0)
                self.assertEqual(records(self.sheet), before)


class QueryMetadataTest(unittest.TestCase):
    def test_search_aliases_and_last_nonempty_value(self):
        for key in ["q", "text", "query", "searchtext", "search", "keyword"]:
            with self.subTest(key=key):
                self.assertEqual(query_from_url(f"https://shop/search?{key}=old&{key}={quote('дрель безударная')}&{key}="), "дрель безударная")
        self.assertEqual(query_from_url("https://shop/?q=old&amp;text=ps+5"), "ps 5")

    def test_decoding_preserves_literal_plus_ampersand_and_percent(self):
        self.assertEqual(query_from_url("https://shop/?q=C%2B%2B+%26+30%25"), "C++ & 30%")
        encoded = quote(quote("ручка гелевая"))
        self.assertEqual(query_from_url("https://shop/?q=" + encoded), "ручка гелевая")
        self.assertIsNone(query_from_url("https://shop/?qid=123&category=tools"))

    def test_url_priority_and_existing_annotation(self):
        page = {"page_url": "https://shop/?q=new", "source_url": "https://shop/?q=old"}
        self.assertEqual(recover_query(page, Path("."))["query"], "new")
        page["query"] = "manual"
        self.assertEqual(recover_query(page, Path("."))["query"], "manual")

    def test_ambiguous_html_links_are_not_search_history(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "page.html").write_text('<div class="pagination"><a href="/?q=old">a</a><a href="/?q=unrelated">b</a></div>', encoding="utf-8")
            page = {"benchmark_html_file": "page.html", "source": "shop", "page_url": "https://shop/"}
            result = recover_query(page, root)
            self.assertIsNone(result["query"])
            self.assertEqual(result["query_source"], "ambiguous_links")

    def test_promotional_link_does_not_become_query(self):
        root = Path(__file__).resolve().parents[1] / "industrial"
        result = recover_query({"benchmark_html_file": "html/officemag.ru-run-none2.html",
                                "source": "officemag.ru", "page_url": "https://www.officemag.ru/search/?q="}, root)
        self.assertIsNone(result["query"])

    def test_metadata_preserves_zero_and_done_product_labels(self):
        book = Workbook()
        gt = book.active
        gt.title = "GroundTruth"
        headers = ["html_file", "query", "is_product", "name_gt", "price_gt", "product_url_gt", "in_stock_gt", "annotator", "annotation_status", "notes"]
        gt.append(headers)
        gt.append(["p.html", None, 1, "Manual product", 0, "https://shop/p", None, "Human", "done", "Original note"])
        gt.append(["empty.html", None, None, None, None, None, None, None, "pending", None])
        pages = book.create_sheet("Pages")
        pages.append(["benchmark_html_file", "source", "query", "page_url", "source_url"])
        pages.append(["p.html", "shop", None, "https://shop/?text=" + quote("дрель"), None])
        pages.append(["empty.html", "shop", None, "https://shop/?q=abcdef", None])
        # Plain metadata recovery must never mark an empty page negative.
        update_metadata(book, Path("."))
        self.assertIsNone(gt.cell(3, 3).value)
        result = update_metadata(book, Path("."), confirm_empty_pending=True, confirm_unpriced_unavailable=True)
        self.assertEqual(result["confirmed_empty_rows"], 1)
        self.assertEqual(result["confirmed_unavailable_rows"], 1)
        rows = records(gt)
        self.assertEqual(rows[0]["price_gt"], 0)
        self.assertEqual(rows[0]["name_gt"], "Manual product")
        self.assertEqual(rows[0]["annotator"], "Human")
        self.assertEqual(rows[0]["query"], "дрель")
        self.assertEqual(rows[0]["in_stock_gt"], 0)
        self.assertEqual(rows[1]["is_product"], 0)
        self.assertEqual(rows[1]["annotation_status"], "done")
        before = list(gt.values)
        again = update_metadata(book, Path("."), confirm_empty_pending=True, confirm_unpriced_unavailable=True)
        self.assertTrue(all(again[key] == 0 for key in ["query_cells_filled", "confirmed_empty_rows", "confirmed_unavailable_rows", "audit_cells_updated"]))
        self.assertEqual(list(gt.values), before)


if __name__ == "__main__":
    unittest.main()
