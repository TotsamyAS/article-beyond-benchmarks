"""Review exports preserve GT identity and literal values across methods."""
import tempfile
from pathlib import Path
import unittest

import pandas as pd
from openpyxl import load_workbook

from tests.test_baseline_determined import load_benchmark


class ReviewWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ns = load_benchmark()

    def fixture(self):
        common = dict(dataset="demo", entity_id="entity-1", gt_id=1, page_id="page-1", source="shop.test",
                      query="дрель", html_file="html/a.html", product_url="https://shop.test/p/1",
                      gt_name="=Название товара", gt_price=None, gt_price_state="absent", name_similarity=1,
                      notes="", error_labels="", price_presence_ok=True)
        rows = [dict(common, variant="b1", benchmark_run_id="run-1", pred_name=None, pred_price=None,
                     match_status="missing_prediction", name_ok=False, price_ok=None, name_TP=0),
                dict(common, variant="b2", benchmark_run_id="run-2", pred_name="=Название товара", pred_price=0,
                     match_status="matched_gt", name_ok=True, price_ok=False, name_TP=1),
                dict(common, entity_id=None, gt_id=None, gt_name=None, gt_price_state=None,
                     product_url="https://shop.test/p/2", variant="b2", benchmark_run_id="run-2",
                     pred_name="Лишний товар", pred_price=100, match_status="unmatched_prediction",
                     name_ok=False, price_ok=False, name_TP=0)]
        return pd.DataFrame(rows)

    def test_entities_side_by_side_extras_separate_literal_excel_values(self):
        evaluation = self.fixture()
        comparison, extras = self.ns["build_review_comparison"](evaluation)
        self.assertEqual(len(comparison), 1)
        self.assertEqual(len(extras), 1)
        self.assertEqual(comparison.iloc[0]["b1__match_status"], "missing_prediction")
        self.assertEqual(comparison.iloc[0]["b2__pred_price"], 0)
        metrics = pd.DataFrame([dict(dataset="demo", variant=v) for v in ["b1", "b2"]])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.xlsx"
            self.ns["export_review_workbook"](path, evaluation, metrics)
            workbook = load_workbook(path)
            sheet = workbook["Сравнение"]
            columns = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, columns["gt_name"]).value, "=Название товара")
            self.assertEqual(sheet.cell(2, columns["gt_name"]).data_type, "s")
            self.assertEqual(sheet.cell(2, columns["b2__pred_price"]).value, 0)
            self.assertIsNone(sheet.cell(2, columns["gt_price"]).value)
            self.assertFalse(sheet.cell(2, columns["b1__name_ok"]).value)
            self.assertEqual(workbook["Лишние"].max_row, 2)
            self.assertEqual(workbook["Все_результаты"].max_row, 4)
            workbook.close()

    def test_ambiguous_entities_or_conflicting_gt_fail_instead_of_merging(self):
        frame = self.fixture()
        with self.assertRaisesRegex(ValueError, "unique GT"):
            self.ns["build_review_comparison"](pd.concat([frame, frame.iloc[[0]]]))
        frame.loc[1, "gt_name"] = "Другой эталон"
        with self.assertRaisesRegex(ValueError, "Conflicting GT"):
            self.ns["build_review_comparison"](frame)

    def test_empty_unscored_run_still_exports_readable_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "empty.xlsx"
            self.ns["export_review_workbook"](path, pd.DataFrame(), pd.DataFrame())
            workbook = load_workbook(path)
            self.assertIn("Сравнение", workbook.sheetnames)
            self.assertEqual(workbook["Сравнение"].max_row, 1)
            workbook.close()

    def test_csv_blanks_and_numeric_ids_match_fresh_results(self):
        frame = self.fixture().astype({"gt_id": object, "query": object})
        frame.loc[0, "query"] = float("nan")
        frame.loc[1, "query"] = ""
        frame.loc[0, "gt_id"] = 1.0
        frame.loc[1, "gt_id"] = "1"
        comparison, _ = self.ns["build_review_comparison"](frame)
        self.assertEqual(len(comparison), 1)
        self.assertEqual(comparison.iloc[0].gt_id, "1")
        frame.loc[1, "gt_id"] = "001"
        with self.assertRaisesRegex(ValueError, "Conflicting GT"):
            self.ns["build_review_comparison"](frame)


if __name__ == "__main__":
    unittest.main()
